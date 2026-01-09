"""
Simple Excel Generator (Non-AI)
Converts geometric table data to Excel.
"""

import io
from typing import List, Dict, Any
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, Border, Side, PatternFill

def generate_excel_geometric(data: Dict[str, Any]) -> io.BytesIO:
    """
    Generate Excel from geometric reconstruction data.
    """
    wb = Workbook()
    if 'Sheet' in wb.sheetnames:
        del wb['Sheet']
        
    pages = data.get('pages', [])
    
    # Check if we have any tables
    has_tables = False
    for page in pages:
        if page.get('tables'):
            has_tables = True
            break
            
    if not has_tables:
        ws = wb.create_sheet("No Data")
        ws['A1'] = "No tables detected in the document."
        output = io.BytesIO()
        wb.save(output)
        output.seek(0)
        return output

    # Create a sheet for each table found
    table_count = 0
    for page in pages:
        page_num = page.get('page_num')
        tables = page.get('tables', [])
        
        for i, table in enumerate(tables):
            table_count += 1
            sheet_name = f"Page {page_num} - Table {i+1}"
            _create_sheet(wb, sheet_name, table)

    output = io.BytesIO()
    wb.save(output)
    output.seek(0)
    return output

def _create_sheet(wb: Workbook, sheet_name: str, table: Dict[str, Any]):
    """Create a formatted sheet for a single table."""
    ws = wb.create_sheet(sheet_name)
    
    rows = table.get('rows', [])
    if not rows:
        return

    # Styles
    header_font = Font(bold=True, color="FFFFFF")
    header_fill = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
    border = Border(left=Side(style='thin'), right=Side(style='thin'), top=Side(style='thin'), bottom=Side(style='thin'))
    
    # We assume first row is header? Geometric doesn't know.
    # But usually first row is header. Let's format it.
    
    # Layout rows
    for r_idx, row_data in enumerate(rows, start=1):
        for c_idx, cell_value in enumerate(row_data, start=1):
            cell = ws.cell(row=r_idx, column=c_idx, value=cell_value)
            cell.border = border
            
            # Header formatting for first row
            if r_idx == 1:
                cell.font = header_font
                cell.fill = header_fill
                
    # Auto-adjust column widths
    for col in ws.columns:
        max_length = 0
        column = col[0].column_letter # Get the column name
        for cell in col:
            try:
                if len(str(cell.value)) > max_length:
                    max_length = len(str(cell.value))
            except:
                pass
        adjusted_width = (max_length + 2)
        ws.column_dimensions[column].width = min(adjusted_width, 50)
